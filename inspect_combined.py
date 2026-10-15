import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation('C:/Users/LAKSHMI/FSD/Combined Review.pptx')

for i in range(len(prs.slides)):
    s = prs.slides[i]
    print(f"\n=== SLIDE {i+1} (shapes: {len(s.shapes)}) ===")
    for idx, sh in enumerate(s.shapes):
        t = sh.text.strip().replace('\n', ' ') if sh.has_text_frame else ''
        print(f"  Shape {idx} ({sh.name}): type={sh.shape_type}, left={sh.left}, top={sh.top}, w={sh.width}, h={sh.height}, text=\"{t[:60]}\"")

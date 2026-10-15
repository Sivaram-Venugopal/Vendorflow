import pptx
from pptx.util import Inches

prs_path = 'C:/Users/LAKSHMI/FSD/VendorFlow_Combined_Review.pptx'
save_path = 'C:/Users/LAKSHMI/FSD/VendorFlow_Combined_Review_v2.pptx'
prs = pptx.Presentation(prs_path)

s3 = prs.slides[2]

# Remove any existing picture shape on Slide 3
for sh in list(s3.shapes):
    if sh.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        el = sh._element
        el.getparent().remove(el)

# Add the new widescreen architecture diagram (11.0 in x 4.85 in)
arch_img = 'C:/Users/LAKSHMI/FSD/docs/images/figure_4_1_architecture.png'
s3.shapes.add_picture(arch_img, Inches(1.15), Inches(1.5), Inches(11.0), Inches(4.85))

prs.save(save_path)
print(f"Updated Slide 3 and saved successfully to {save_path}!")

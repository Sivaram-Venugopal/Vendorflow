import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation('C:/Users/LAKSHMI/FSD/Combined Review.pptx')

for idx in [5, 7, 11, 14, 16]:
    s = prs.slides[idx]
    print(f"\n=== SLIDE {idx+1} ===")
    for i, sh in enumerate(s.shapes):
        t = sh.text.strip().replace('\n', ' ') if sh.has_text_frame else ''
        print(f"  Shape {i}: {sh.name}, type={sh.shape_type}, text=\"{t[:70]}\"")
        if sh.has_table:
            tb = sh.table
            print(f"    Table size: {len(tb.rows)}x{len(tb.columns)}")
            for r in tb.rows:
                row_str = " | ".join([c.text.strip().replace('\n', ' ') for c in r.cells])
                print(f"      {row_str[:80]}")

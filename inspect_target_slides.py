import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation('C:/Users/LAKSHMI/FSD/Combined Review.pptx')

for s_idx in [0, 2, 3, 4, 7, 13]:
    s = prs.slides[s_idx]
    print(f"\n=== SLIDE {s_idx+1} ===")
    for idx, sh in enumerate(s.shapes):
        t = sh.text.strip().replace('\n', ' ') if sh.has_text_frame else ''
        print(f"  Shape {idx} ({sh.name}): left={sh.left}, top={sh.top}, w={sh.width}, h={sh.height}, text=\"{t[:60]}\"")

import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation('C:/Users/LAKSHMI/FSD/0th review - 310624205263.pptx')

print(f"Presentation width: {prs.slide_width}, height: {prs.slide_height}")
for i, slide in enumerate(prs.slides):
    print(f"\n=== SLIDE {i+1} ===")
    for j, shape in enumerate(slide.shapes):
        has_tf = shape.has_text_frame
        txt = shape.text.strip().replace('\n', ' ') if has_tf else ''
        img_info = ' [Picture]' if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE else ''
        print(f"  Shape {j} ({shape.name}): left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}{img_info}")
        if has_tf:
            print(f"    Text: {txt[:100]}...")
